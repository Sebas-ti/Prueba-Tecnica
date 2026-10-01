"""Carga de documentos a texto plano con metadatos de página/sección."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

from app.core.errors import BadRequestError, UnsupportedFileError


@dataclass
class LoadedSection:
    text: str
    page: int | None = None


def load_bytes(filename: str, content: bytes) -> list[LoadedSection]:
    ext = Path(filename).suffix.lower()
    if ext in {".md", ".txt"}:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        return [LoadedSection(text=text)]
    if ext == ".pdf":
        return _load_pdf(content)
    if ext == ".docx":
        return _load_docx(content)
    raise UnsupportedFileError(f"Extensión no soportada: {ext}")


def _load_pdf(content: bytes) -> list[LoadedSection]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise BadRequestError("No se pudo leer el PDF (¿archivo corrupto o cifrado?)") from exc
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            # Reflujo: el PDF corta líneas a mitad de oración; se unen las líneas
            # que no terminan en puntuación y se reparan palabras con guion.
            text = re.sub(r"-\n(?=\w)", "", text)
            text = re.sub(r"(?<![.:!?])\n(?!\n)", " ", text)
            sections.append(LoadedSection(text=text, page=i))
    if not sections:
        # PDF escaneado: en producción se enviaría a Azure AI Document Intelligence (OCR).
        raise BadRequestError("El PDF no contiene texto extraíble (posible documento escaneado)")
    return sections


def _load_docx(content: bytes) -> list[LoadedSection]:
    import docx

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise BadRequestError("No se pudo leer el DOCX") from exc
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style.startswith("título") or style.startswith("titulo"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * int(level)} {text}")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [LoadedSection(text="\n\n".join(lines))]
