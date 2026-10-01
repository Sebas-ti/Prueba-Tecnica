"""Genera los documentos de ejemplo en formato DOCX y PDF.

Sirven para demostrar que la ingesta soporta formatos binarios además de
Markdown. Uso: python -m scripts.generate_sample_docs
"""
from __future__ import annotations

from pathlib import Path

DOCS = Path(__file__).resolve().parents[1] / "data" / "docs"

FAQ = [
    ("¿Cómo registro una solicitud?",
     "Ingrese a GESOL con su usuario de red, seleccione Nueva solicitud, elija la categoría y adjunte la evidencia. "
     "El sistema asigna un número con formato SOL seguido de cuatro dígitos."),
    ("¿Cómo restablezco mi contraseña de GESOL?",
     "GESOL usa la contraseña del Directorio Activo. Para restablecerla use el portal de autoservicio de contraseñas "
     "o llame a la Mesa de Servicios a la extensión 4500."),
    ("¿Puedo cancelar una solicitud?",
     "Sí. El solicitante puede cancelar una solicitud mientras esté en estado Registrada o En análisis. "
     "Después de Aprobada, la cancelación debe pedirse a la Mesa de Servicios."),
    ("¿Cuál es el tamaño máximo de un adjunto?",
     "El tamaño máximo por adjunto es de 25 MB y se permiten hasta 10 adjuntos por solicitud. "
     "Los formatos permitidos son PDF, imágenes, documentos de Office y archivos comprimidos ZIP."),
    ("¿Quién aprueba mis solicitudes?",
     "Las solicitudes son aprobadas por el jefe inmediato registrado en SAP HCM y por el dueño del servicio. "
     "Si el aprobador está ausente, la solicitud se delega automáticamente después de 3 días hábiles."),
]

ACTA = """Acta del Comité de Arquitectura No. 2026-09

Fecha: 15 de septiembre de 2026. Participantes: Arquitectura Empresarial, Seguridad de la Información, Aplicaciones Core, Infraestructura Cloud y la Gerencia de Servicios Compartidos.

Tema 1. Asistente inteligente para GESOL. El comité aprobó la construcción de un asistente basado en RAG sobre Azure como primer incremento de la modernización de GESOL. El asistente debe operar en la suscripción corporativa, usar Azure OpenAI con acceso privado, Azure AI Search como índice y registrar todas las interacciones para auditoría.

Tema 2. Criterios de salida a producción. El asistente solo podrá salir a producción si en la evaluación supera un 85 por ciento de respuestas correctas, no responde preguntas sin información en las fuentes y supera las pruebas de inyección de instrucciones definidas por Seguridad.

Tema 3. Presupuesto. Se asignó un presupuesto mensual máximo de 1.500 dólares para el consumo de servicios de IA del piloto. Se deben configurar alertas de costo al 80 por ciento del presupuesto.

Tema 4. Decisión sobre la base de datos. Se decidió mantener Oracle durante el piloto y evaluar la migración a Azure Database for PostgreSQL en 2027.

Compromisos: Arquitectura presentará el diseño detallado el 30 de septiembre de 2026. Seguridad entregará el set de pruebas de inyección el 10 de octubre de 2026."""


def build_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("Preguntas Frecuentes de la Mesa de Servicios sobre GESOL", level=1)
    d.add_paragraph("Documento de apoyo para usuarios finales. Versión 2026.")
    for q, a in FAQ:
        d.add_heading(q, level=2)
        d.add_paragraph(a)
    d.save(path)


def build_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    paragraphs = ACTA.split("\n\n")
    story.append(Paragraph(paragraphs[0], styles["Title"]))
    for p in paragraphs[1:]:
        story.append(Paragraph(p, styles["BodyText"]))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(str(path), pagesize=letter, title="Acta Comité de Arquitectura 2026-09").build(story)


if __name__ == "__main__":
    DOCS.mkdir(parents=True, exist_ok=True)
    build_docx(DOCS / "faq_mesa_servicios.docx")
    build_pdf(DOCS / "acta_comite_arquitectura_2026_09.pdf")
    print("Documentos generados en", DOCS)
